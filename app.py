# -*- coding: utf-8 -*-
"""
KPI ENTREGA DE INFORMES — Tablero de control y gestión de la matriz
ADEMINSAC | Gestión de Proyectos · Mejora Continua

KPIs principales
- OTD  : % de informes entregados dentro del plazo (On-Time Delivery)
- Lead Time : días desde la ejecución del servicio hasta la entrega del informe
- FPY  : % de informes aprobados sin revisiones (First Pass Yield)
- Backlog / Vencidos : informes pendientes y pendientes fuera de plazo
- Estabilidad del proceso : gráfico de control I-MR del lead time
"""
from __future__ import annotations

import base64
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

# ════════════════════════════════════════════════════════════════════
# CONFIGURACIÓN GENERAL
# ════════════════════════════════════════════════════════════════════
st.set_page_config(
    page_title="KPI Entrega de Informes | ADEMINSAC",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

DATA_DIR = Path(__file__).parent / "data"
DATA_FILE = DATA_DIR / "matriz_kpi_informes.xlsx"
SHEET = "MATRIZ"
TZ = ZoneInfo("America/Lima")

COLUMNS = [
    "ID", "N° OT", "Especialidad", "Tipo de Informe", "Unidad / Área",
    "Equipo / TAG", "Responsable", "Fecha Ejecución", "Fecha Compromiso",
    "Fecha Entrega", "Estado", "N° Revisiones", "Causa de Retraso", "Observaciones",
]
DATE_COLS = ["Fecha Ejecución", "Fecha Compromiso", "Fecha Entrega"]
TEXT_COLS = [c for c in COLUMNS if c not in DATE_COLS + ["N° Revisiones"]]

ESPECIALIDADES = ["END", "Estáticos", "Dinámicos", "Instrumentación", "Electricidad", "SSOMA"]
TIPOS = ["Informe Técnico", "Informe END", "Informe de Hallazgos", "Informe de Calibración",
         "Informe SSOMA", "Informe Final"]
ESTADOS = ["Pendiente", "En elaboración", "En revisión", "Entregado", "Observado", "Aprobado", "Anulado"]
CAUSAS = ["Carga de trabajo", "Demora en revisión interna", "Falta de información de campo",
          "Espera de resultados / laboratorio", "Observaciones del cliente", "Acceso / permisos",
          "Cambio de alcance", "Otros"]

# Paleta apta para daltonismo (azul/naranja como par principal)
AZUL, AZUL_CL, NARANJA, GRIS, ROJO = "#1F4E79", "#5DADE2", "#E67E22", "#A6ACAF", "#B03A2E"
COLOR_SIT = {
    "Entregado a tiempo": AZUL, "Entregado fuera de plazo": NARANJA,
    "Entregado (sin plazo)": GRIS, "Pendiente en plazo": AZUL_CL, "Pendiente vencido": ROJO,
}

ALIASES = {
    "ID": ["id", "item", "codigo", "cod", "correlativo", "codinforme", "idinforme", "ninforme", "nroinforme"],
    "N° OT": ["not", "ot", "nroot", "numeroot", "ordendetrabajo", "ordentrabajo", "nordentrabajo"],
    "Especialidad": ["especialidad", "disciplina", "areatecnica", "especialidadtecnica"],
    "Tipo de Informe": ["tipodeinforme", "tipoinforme", "tipodocumento", "tipo"],
    "Unidad / Área": ["unidadarea", "unidad", "area", "planta", "ubicacion"],
    "Equipo / TAG": ["equipotag", "tag", "equipo", "activo"],
    "Responsable": ["responsable", "inspector", "elaboradopor", "autor", "encargado"],
    "Fecha Ejecución": ["fechaejecucion", "fechainspeccion", "fechaservicio", "fechacampo",
                        "fechatrabajo", "fechaejec", "fechadeejecucion", "fechadeinspeccion"],
    "Fecha Compromiso": ["fechacompromiso", "fechalimite", "fechaplazo", "fechaprogramada",
                         "fechaobjetivo", "fechadecompromiso", "fechalimitedeentrega"],
    "Fecha Entrega": ["fechaentrega", "fechadeentrega", "fechaenvio", "fechaemision", "fechareal"],
    "Estado": ["estado", "status", "situacion", "estadoinforme"],
    "N° Revisiones": ["nrevisiones", "revisiones", "nrorevisiones", "rev", "reprocesos", "nrodeobservaciones"],
    "Causa de Retraso": ["causaderetraso", "causaretraso", "motivoderetraso", "motivoretraso", "causa"],
    "Observaciones": ["observaciones", "comentarios", "obs", "notas"],
}

st.markdown(
    """
    <style>
    [data-testid="stMetric"]{background:rgba(31,78,121,.06);border:1px solid rgba(31,78,121,.18);
        border-radius:10px;padding:10px 14px;}
    [data-testid="stMetricLabel"] p{font-weight:600;}
    .hdr{padding:.4rem 0 .2rem 0;border-bottom:3px solid #1F4E79;margin-bottom:.8rem;}
    .hdr h1{font-size:1.65rem;margin:0;color:#1F4E79;}
    .hdr p{margin:.1rem 0 0 0;opacity:.75;}
    </style>
    """,
    unsafe_allow_html=True,
)


# ════════════════════════════════════════════════════════════════════
# UTILIDADES: FECHAS Y FERIADOS (PERÚ)
# ════════════════════════════════════════════════════════════════════
def hoy_lima() -> date:
    return datetime.now(TZ).date()


def _pascua(y: int) -> date:
    a, b, c = y % 19, y // 100, y % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return date(y, mes, dia)


@st.cache_data
def feriados_peru(y0: int, y1: int) -> np.ndarray:
    """Feriados nacionales del Perú (fijos + Jueves y Viernes Santo). Editar si cambia la normativa."""
    fijos = [(1, 1), (5, 1), (6, 7), (6, 29), (7, 23), (7, 28), (7, 29), (8, 6),
             (8, 30), (10, 8), (11, 1), (12, 8), (12, 9), (12, 25)]
    dias = []
    for y in range(y0, y1 + 1):
        dias += [date(y, m, d) for m, d in fijos]
        p = _pascua(y)
        dias += [p - timedelta(days=3), p - timedelta(days=2)]
    return np.array(dias, dtype="datetime64[D]")


HOL = feriados_peru(2018, hoy_lima().year + 2)


def dias_entre(a: pd.Series, b: pd.Series, habiles: bool) -> pd.Series:
    """Días de b respecto de a (positivo si b es posterior)."""
    out = pd.Series(np.nan, index=a.index, dtype="float")
    m = a.notna() & b.notna()
    if m.any():
        if habiles:
            out[m] = np.busday_count(a[m].values.astype("datetime64[D]"),
                                     b[m].values.astype("datetime64[D]"), holidays=HOL)
        else:
            out[m] = (b[m].dt.normalize() - a[m].dt.normalize()).dt.days
    return out


def sumar_dias(s: pd.Series, n: int, habiles: bool) -> pd.Series:
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    m = s.notna()
    if m.any():
        if habiles:
            r = np.busday_offset(s[m].values.astype("datetime64[D]"), n, roll="forward", holidays=HOL)
            out[m] = pd.to_datetime(r)
        else:
            out[m] = s[m].dt.normalize() + pd.Timedelta(days=n)
    return out


def a_fecha(s: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(s):
        return pd.to_datetime(s).dt.tz_localize(None) if getattr(s.dt, "tz", None) else s
    num = pd.to_numeric(s, errors="coerce")
    out = pd.to_datetime(s.astype(str).where(s.notna(), None), errors="coerce",
                         dayfirst=True, format="mixed")
    serial = num.notna() & num.between(20000, 80000)          # fechas seriales de Excel
    if serial.any():
        out[serial] = pd.to_datetime(num[serial], unit="D", origin="1899-12-30")
    return out


# ════════════════════════════════════════════════════════════════════
# UTILIDADES: ESTANDARIZACIÓN DE LA MATRIZ
# ════════════════════════════════════════════════════════════════════
def _norm(txt) -> str:
    t = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", t)


def mapeo_automatico(columnas) -> dict:
    """Devuelve {columna_estándar: columna_origen} reconociendo nombres equivalentes."""
    norm = {c: _norm(c) for c in columnas}
    usados, mapa = set(), {}
    for std in COLUMNS:                                   # 1) coincidencia exacta
        for c, n in norm.items():
            if c not in usados and (n == _norm(std) or n in ALIASES[std]):
                mapa[std] = c
                usados.add(c)
                break
    for std in COLUMNS:                                   # 2) coincidencia parcial
        if std in mapa:
            continue
        for c, n in norm.items():
            if c not in usados and any(len(a) >= 5 and a in n for a in ALIASES[std]):
                mapa[std] = c
                usados.add(c)
                break
    return mapa


def siguiente_id(df: pd.DataFrame, k: int = 1) -> list[str]:
    nums = df["ID"].astype(str).str.extract(r"(\d+)\s*$")[0].dropna()
    base = int(pd.to_numeric(nums, errors="coerce").max()) if len(nums) else 0
    return [f"INF-{base + i:04d}" for i in range(1, k + 1)]


def estandarizar(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[COLUMNS]
    for c in DATE_COLS:
        df[c] = a_fecha(df[c])
    df["N° Revisiones"] = pd.to_numeric(df["N° Revisiones"], errors="coerce").fillna(0).clip(lower=0).astype(int)
    for c in TEXT_COLS:
        df[c] = (df[c].astype(object).where(df[c].notna(), "").astype(str).str.strip()
                 .replace({"nan": "", "None": "", "NaT": "", "<NA>": ""}))
    # quitar filas totalmente vacías
    vacia = (df[TEXT_COLS].eq("").all(axis=1)) & df[DATE_COLS].isna().all(axis=1)
    df = df[~vacia].reset_index(drop=True)
    # estado: homologar mayúsculas / completar según fecha de entrega
    homolog = {_norm(e): e for e in ESTADOS}
    df["Estado"] = df["Estado"].map(lambda x: homolog.get(_norm(x), x))
    sin_estado = df["Estado"].eq("")
    df.loc[sin_estado, "Estado"] = np.where(df.loc[sin_estado, "Fecha Entrega"].notna(), "Entregado", "Pendiente")
    # IDs faltantes
    falt = df["ID"].eq("")
    if falt.any():
        df.loc[falt, "ID"] = siguiente_id(df, int(falt.sum()))
    return df


def desde_crudo(raw: pd.DataFrame, mapa: dict | None = None) -> pd.DataFrame:
    raw = raw.dropna(how="all").dropna(axis=1, how="all")
    raw.columns = [str(c).strip() for c in raw.columns]
    mapa = mapa if mapa is not None else mapeo_automatico(raw.columns)
    out = pd.DataFrame({std: raw[src] for std, src in mapa.items() if src in raw.columns})
    return estandarizar(out)


def opciones(df: pd.DataFrame, col: str, base: list[str]) -> list[str]:
    extra = [v for v in df[col].dropna().unique() if v and v not in base]
    return base + sorted(extra)


# ════════════════════════════════════════════════════════════════════
# PERSISTENCIA (archivo local + GitHub opcional)
# ════════════════════════════════════════════════════════════════════
def _secret(*keys):
    try:
        v = st.secrets
        for k in keys:
            v = v[k]
        return v
    except Exception:
        return None


def a_excel(df: pd.DataFrame, extras: dict[str, pd.DataFrame] | None = None) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY", date_format="DD/MM/YYYY") as w:
        hojas = {SHEET: df, **(extras or {})}
        for nombre, data in hojas.items():
            data.to_excel(w, sheet_name=nombre, index=False)
            ws = w.sheets[nombre]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(data.columns, start=1):
                largo = max([len(str(col))] + [len(str(v)) for v in data[col].head(300)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, largo + 2), 45)
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E79")
    return buf.getvalue()


def subir_github(contenido: bytes, mensaje: str) -> tuple[bool, str]:
    token, repo = _secret("github", "token"), _secret("github", "repo")
    if not token or not repo:
        return False, ""
    import requests
    rama = _secret("github", "branch") or "main"
    ruta = _secret("github", "path") or "data/matriz_kpi_informes.xlsx"
    url = f"https://api.github.com/repos/{repo}/contents/{ruta}"
    hdr = {"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json"}
    try:
        r = requests.get(url, headers=hdr, params={"ref": rama}, timeout=20)
        sha = r.json().get("sha") if r.status_code == 200 else None
        payload = {"message": mensaje, "content": base64.b64encode(contenido).decode(), "branch": rama}
        if sha:
            payload["sha"] = sha
        r = requests.put(url, headers=hdr, json=payload, timeout=30)
        return r.status_code in (200, 201), f"GitHub respondió {r.status_code}"
    except Exception as e:  # noqa: BLE001
        return False, f"GitHub: {e}"


def guardar(df: pd.DataFrame, motivo: str) -> None:
    df = estandarizar(df)
    st.session_state.df = df
    st.session_state.origen = "archivo"
    contenido = a_excel(df)
    msgs = []
    try:
        DATA_DIR.mkdir(exist_ok=True)
        DATA_FILE.write_bytes(contenido)
        msgs.append("✅ Matriz guardada en el servidor")
    except Exception as e:  # noqa: BLE001
        msgs.append(f"⚠️ No se pudo escribir el archivo local: {e}")
    ok, info = subir_github(contenido, f"{motivo} · {datetime.now(TZ):%d/%m/%Y %H:%M}")
    if ok:
        msgs.append("✅ Respaldo sincronizado en GitHub")
    elif info:
        msgs.append(f"⚠️ {info}")
    st.session_state.ultimo_guardado = datetime.now(TZ)
    st.session_state.flash = f"**{motivo}** · " + " · ".join(msgs)


# ════════════════════════════════════════════════════════════════════
# DATOS DE DEMOSTRACIÓN (solo si no existe matriz)
# ════════════════════════════════════════════════════════════════════
def datos_demo(n: int = 240, seed: int = 11) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    hoy = pd.Timestamp(hoy_lima())
    inicio = hoy - pd.Timedelta(days=270)
    esp_p = {"END": .30, "Estáticos": .24, "Dinámicos": .15, "Instrumentación": .15, "Electricidad": .10, "SSOMA": .06}
    tipo_por_esp = {"END": "Informe END", "Estáticos": "Informe Técnico", "Dinámicos": "Informe Técnico",
                    "Instrumentación": "Informe de Calibración", "Electricidad": "Informe Técnico", "SSOMA": "Informe SSOMA"}
    lt_base = {"END": 4.0, "Estáticos": 5.5, "Dinámicos": 4.5, "Instrumentación": 3.5, "Electricidad": 4.0, "SSOMA": 3.0}
    unidades = ["U-02 Destilación", "U-21", "U-23", "Conversión", "Servicios Industriales", "Tanques y Terminal"]
    causas_p = [.30, .22, .18, .12, .10, .05, .02, .01]
    filas = []
    for i in range(n):
        esp = rng.choice(list(esp_p), p=list(esp_p.values()))
        ejec = np.busday_offset((inicio + pd.Timedelta(days=int(rng.integers(0, 268)))).date(), 0,
                                roll="forward", holidays=HOL)
        avance = (pd.Timestamp(ejec) - inicio).days / 270          # mejora gradual en el tiempo
        lt = max(1, int(round(rng.gamma(4, lt_base[esp] / 4) * (1.15 - .35 * avance))))
        entrega = pd.Timestamp(np.busday_offset(ejec, lt, holidays=HOL))
        rev = int(rng.choice([0, 1, 2, 3], p=[.64, .23, .10, .03]))
        if entrega > hoy or rng.random() < .05:
            f_ent, estado = pd.NaT, rng.choice(["Pendiente", "En elaboración", "En revisión"])
        else:
            f_ent, estado = entrega, ("Observado" if rev > 1 else rng.choice(["Entregado", "Aprobado"], p=[.4, .6]))
        causa = rng.choice(CAUSAS, p=causas_p) if (pd.notna(f_ent) and lt > 5) else ""
        filas.append({
            "ID": f"INF-{i + 1:04d}", "N° OT": f"OT-{4500100 + i}", "Especialidad": esp,
            "Tipo de Informe": tipo_por_esp[esp], "Unidad / Área": rng.choice(unidades),
            "Equipo / TAG": f"{rng.choice(['E', 'P', 'V', 'T', 'PSV'])}-{rng.integers(100, 999)}",
            "Responsable": f"{esp} · Resp. {rng.integers(1, 4)}", "Fecha Ejecución": pd.Timestamp(ejec),
            "Fecha Compromiso": pd.NaT, "Fecha Entrega": f_ent, "Estado": estado,
            "N° Revisiones": rev if pd.notna(f_ent) else 0, "Causa de Retraso": causa, "Observaciones": "",
        })
    return estandarizar(pd.DataFrame(filas))


def cargar_inicial() -> tuple[pd.DataFrame, str]:
    if DATA_FILE.exists():
        try:
            return desde_crudo(pd.read_excel(DATA_FILE, sheet_name=0)), "archivo"
        except Exception:  # noqa: BLE001
            pass
    return datos_demo(), "demo"


# ════════════════════════════════════════════════════════════════════
# CÁLCULO DE INDICADORES
# ════════════════════════════════════════════════════════════════════
def calcular(df: pd.DataFrame, plazo: int, habiles: bool, hoy: date) -> pd.DataFrame:
    d = df[df["Estado"] != "Anulado"].copy()
    lim = d["Fecha Compromiso"].copy()
    sin = lim.isna()
    lim.loc[sin] = sumar_dias(d.loc[sin, "Fecha Ejecución"], plazo, habiles)
    d["Fecha Límite"] = lim
    d["Entregado"] = d["Fecha Entrega"].notna()
    d["Lead Time"] = dias_entre(d["Fecha Ejecución"], d["Fecha Entrega"], habiles)
    d["Desviación"] = dias_entre(d["Fecha Límite"], d["Fecha Entrega"], habiles)
    hoy_s = pd.Series(pd.Timestamp(hoy), index=d.index)
    d["Días Vencido"] = dias_entre(d["Fecha Límite"], hoy_s, habiles).where(~d["Entregado"]).clip(lower=0)
    d["Antigüedad"] = dias_entre(d["Fecha Ejecución"], hoy_s, habiles).where(~d["Entregado"])
    d["A Tiempo"] = (d["Desviación"] <= 0).astype(float).where(d["Entregado"] & d["Fecha Límite"].notna())
    d["Sin Revisión"] = (d["N° Revisiones"] == 0).astype(float).where(d["Entregado"])
    d["Situación"] = np.select(
        [d["Entregado"] & (d["A Tiempo"] == 1), d["Entregado"] & (d["A Tiempo"] == 0), d["Entregado"],
         ~d["Entregado"] & (d["Días Vencido"] > 0)],
        ["Entregado a tiempo", "Entregado fuera de plazo", "Entregado (sin plazo)", "Pendiente vencido"],
        default="Pendiente en plazo",
    )
    d["Mes Entrega"] = d["Fecha Entrega"].dt.strftime("%Y-%m")
    d["Mes Ejecución"] = d["Fecha Ejecución"].dt.strftime("%Y-%m")
    d["Especialidad"] = d["Especialidad"].replace("", "Sin especialidad")
    d["Responsable"] = d["Responsable"].replace("", "Sin responsable")
    d["Unidad / Área"] = d["Unidad / Área"].replace("", "Sin unidad")
    return d


def estilo(fig: go.Figure, h: int = 380, titulo: str | None = None) -> go.Figure:
    fig.update_layout(
        template="plotly_white", height=h, margin=dict(l=10, r=10, t=55 if titulo else 25, b=10),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0, title=None),
        font=dict(family="Segoe UI, Arial, sans-serif", size=12),
        title=dict(text=titulo, font=dict(size=15, color=AZUL)) if titulo else None,
    )
    return fig


def mostrar(fig: go.Figure) -> None:
    st.plotly_chart(fig, use_container_width=True, config={"displaylogo": False})


def sin_datos(msg: str = "No hay datos suficientes para este gráfico con los filtros actuales.") -> None:
    st.info(msg, icon="ℹ️")


def hallazgos(d: pd.DataFrame, meta: float) -> list[str]:
    ent = d[d["A Tiempo"].notna()]
    out = []
    if len(ent):
        otd = ent["A Tiempo"].mean() * 100
        icon = "✅" if otd >= meta else "⚠️"
        out.append(f"{icon} Cumplimiento de plazo global de **{otd:.1f}%** frente a una meta de {meta:.0f}% "
                   f"({'+' if otd >= meta else ''}{otd - meta:.1f} pp).")
        esp = ent.groupby("Especialidad")["A Tiempo"].agg(["mean", "count"])
        esp = esp[esp["count"] >= 3]
        if len(esp) > 1:
            w = esp["mean"].idxmin()
            out.append(f"🎯 La especialidad con menor cumplimiento es **{w}** ({esp.loc[w, 'mean'] * 100:.0f}%, "
                       f"n={int(esp.loc[w, 'count'])}); se recomienda priorizarla en el análisis de causa raíz.")
        meses = ent.groupby("Mes Entrega")["A Tiempo"].mean().sort_index() * 100
        if len(meses) >= 6:
            ult, ant = meses.iloc[-3:].mean(), meses.iloc[-6:-3].mean()
            tend = "mejora" if ult > ant else "deterioro"
            out.append(f"📈 Tendencia trimestral: **{tend}** del OTD ({ant:.0f}% → {ult:.0f}%) "
                       f"comparando los últimos 3 meses con los 3 anteriores.")
        tarde = ent[ent["A Tiempo"] == 0]
        if len(tarde):
            c = tarde["Causa de Retraso"].replace("", "Sin causa registrada").value_counts(normalize=True)
            out.append(f"🔍 Principal causa de retraso: **{c.index[0]}** ({c.iloc[0] * 100:.0f}% de los informes "
                       f"fuera de plazo).")
            if (tarde["Causa de Retraso"] == "").mean() > .2:
                out.append("📝 Más del 20% de los retrasos no tiene causa registrada; completar este campo "
                           "fortalece el análisis de Pareto.")
        fpy = d["Sin Revisión"].mean() * 100
        if not np.isnan(fpy):
            out.append(f"🧾 El {fpy:.0f}% de los informes se aprueba sin revisiones (FPY).")
    pend = d[~d["Entregado"]]
    venc = pend[pend["Días Vencido"] > 0]
    if len(venc):
        out.append(f"⏰ **{len(venc)}** informe(s) pendiente(s) vencido(s); el más antiguo acumula "
                   f"{int(venc['Días Vencido'].max())} días de atraso.")
    elif len(pend):
        out.append(f"🟦 {len(pend)} informe(s) pendiente(s), todos dentro del plazo.")
    return out


# ════════════════════════════════════════════════════════════════════
# COMPONENTES DE GESTIÓN DE LA MATRIZ
# ════════════════════════════════════════════════════════════════════
def form_registro(prefix: str) -> None:
    if not puede_editar():
        st.warning("🔒 Ingrese la clave de edición en la barra lateral para registrar informes.")
        return
    df = st.session_state.df
    k = lambda n: f"{prefix}_{n}"  # noqa: E731
    resp_exist = sorted(v for v in df["Responsable"].unique() if v)
    with st.form(k("form"), clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        ot = c1.text_input("N° OT", key=k("ot"))
        esp = c2.selectbox("Especialidad *", opciones(df, "Especialidad", ESPECIALIDADES), key=k("esp"))
        tipo = c3.selectbox("Tipo de Informe", opciones(df, "Tipo de Informe", TIPOS), key=k("tipo"))
        c4, c5, c6 = st.columns(3)
        unidad = c4.text_input("Unidad / Área", key=k("uni"))
        tag = c5.text_input("Equipo / TAG", key=k("tag"))
        resp_sel = c6.selectbox("Responsable", ["—"] + resp_exist, key=k("resp"))
        c7, c8, c9 = st.columns(3)
        f_ej = c7.date_input("Fecha Ejecución *", value=hoy_lima(), format="DD/MM/YYYY", key=k("fej"))
        f_co = c8.date_input("Fecha Compromiso (opcional)", value=None, format="DD/MM/YYYY", key=k("fco"),
                             help="Si se deja vacía se calcula con el plazo estándar configurado.")
        f_en = c9.date_input("Fecha Entrega", value=None, format="DD/MM/YYYY", key=k("fen"))
        c10, c11, c12 = st.columns(3)
        estado = c10.selectbox("Estado", ESTADOS, key=k("est"))
        rev = c11.number_input("N° Revisiones", min_value=0, max_value=30, value=0, step=1, key=k("rev"))
        causa = c12.selectbox("Causa de Retraso", [""] + opciones(df, "Causa de Retraso", CAUSAS), key=k("cau"))
        c13, c14 = st.columns([1, 2])
        resp_new = c13.text_input("…o registrar nuevo responsable", key=k("respn"))
        obs = c14.text_input("Observaciones", key=k("obs"))
        enviar = st.form_submit_button("💾 Registrar informe", type="primary")
    if enviar:
        errores = []
        if f_en and f_ej and f_en < f_ej:
            errores.append("La fecha de entrega no puede ser anterior a la fecha de ejecución.")
        if estado in ("Entregado", "Observado", "Aprobado") and not f_en:
            errores.append(f"El estado «{estado}» requiere registrar la Fecha Entrega.")
        if errores:
            for e in errores:
                st.error(e)
            return
        if f_en and estado in ("Pendiente", "En elaboración", "En revisión"):
            estado = "Entregado"
        nuevo_id = siguiente_id(df)[0]
        fila = {
            "ID": nuevo_id, "N° OT": ot, "Especialidad": esp, "Tipo de Informe": tipo, "Unidad / Área": unidad,
            "Equipo / TAG": tag, "Responsable": resp_new.strip() or ("" if resp_sel == "—" else resp_sel),
            "Fecha Ejecución": pd.Timestamp(f_ej) if f_ej else pd.NaT,
            "Fecha Compromiso": pd.Timestamp(f_co) if f_co else pd.NaT,
            "Fecha Entrega": pd.Timestamp(f_en) if f_en else pd.NaT, "Estado": estado,
            "N° Revisiones": int(rev), "Causa de Retraso": causa, "Observaciones": obs,
        }
        guardar(pd.concat([df, pd.DataFrame([fila])], ignore_index=True), f"Registro {nuevo_id} añadido")
        st.rerun()


def carga_matriz(prefix: str, compacto: bool = False) -> None:
    if not puede_editar():
        st.warning("🔒 Ingrese la clave de edición en la barra lateral para subir una matriz.")
        return
    up = st.file_uploader("Matriz actualizada (.xlsx, .xlsm o .csv)", type=["xlsx", "xlsm", "csv"], key=f"{prefix}_up")
    if not up:
        return
    try:
        crudo = up.getvalue()
        c1, c2 = st.columns(2) if not compacto else (st.container(), st.container())
        fila_hdr = c2.number_input("Fila de encabezados", 1, 30, 1, key=f"{prefix}_hdr")
        if up.name.lower().endswith(".csv"):
            raw = pd.read_csv(io.BytesIO(crudo), sep=None, engine="python", header=fila_hdr - 1,
                              encoding_errors="replace")
        else:
            xl = pd.ExcelFile(io.BytesIO(crudo))
            idx = next((i for i, s in enumerate(xl.sheet_names) if "matriz" in s.lower()), 0)
            hoja = c1.selectbox("Hoja", xl.sheet_names, index=idx, key=f"{prefix}_hoja")
            raw = xl.parse(hoja, header=fila_hdr - 1)
        raw = raw.dropna(how="all").dropna(axis=1, how="all")
        raw.columns = [str(c).strip() for c in raw.columns]
    except Exception as e:  # noqa: BLE001
        st.error(f"No se pudo leer el archivo: {e}")
        return

    auto = mapeo_automatico(raw.columns)
    faltan = [c for c in ["Especialidad", "Fecha Ejecución", "Fecha Entrega"] if c not in auto]
    NO = "— (no disponible)"
    mapa = {}
    with st.expander(f"🔗 Mapeo de columnas ({len(auto)}/{len(COLUMNS)} reconocidas)", expanded=bool(faltan)):
        cols = st.columns(2)
        for i, std in enumerate(COLUMNS):
            ops = [NO] + list(raw.columns)
            sel = cols[i % 2].selectbox(std, ops, index=ops.index(auto[std]) if std in auto else 0,
                                        key=f"{prefix}_map_{i}")
            if sel != NO:
                mapa[std] = sel
    if faltan:
        st.warning("Columnas clave no reconocidas automáticamente: " + ", ".join(faltan) +
                   ". Asígnelas en el mapeo para calcular los KPI correctamente.")
    modo = st.radio("Modo de carga", ["Reemplazar matriz completa", "Anexar / actualizar por ID"],
                    key=f"{prefix}_modo", horizontal=not compacto,
                    help="«Anexar» agrega registros nuevos y actualiza los que tengan el mismo ID.")
    nueva = desde_crudo(raw, mapa)
    st.caption(f"Vista previa: {len(nueva)} registros válidos")
    st.dataframe(nueva.head(8), use_container_width=True, hide_index=True)
    if st.button("✅ Confirmar carga", type="primary", key=f"{prefix}_ok"):
        if modo.startswith("Reemplazar"):
            final, motivo = nueva, f"Matriz reemplazada ({len(nueva)} registros)"
        else:
            final = pd.concat([st.session_state.df, nueva], ignore_index=True)
            final = final.drop_duplicates(subset="ID", keep="last")
            motivo = f"Matriz anexada/actualizada ({len(nueva)} registros procesados)"
        guardar(final, motivo)
        st.rerun()


def editor_matriz() -> None:
    if not puede_editar():
        st.warning("🔒 Ingrese la clave de edición en la barra lateral para modificar la matriz.")
        st.dataframe(st.session_state.df, use_container_width=True, hide_index=True)
        return
    df = st.session_state.df
    st.caption("Edite directamente las celdas. Para agregar filas use la última fila vacía; para eliminar, "
               "seleccione la fila (casilla izquierda) y presione Supr. Luego pulse **Guardar cambios**.")
    cfg = {
        "ID": st.column_config.TextColumn("ID", help="Se autogenera si se deja vacío"),
        "Especialidad": st.column_config.SelectboxColumn(options=opciones(df, "Especialidad", ESPECIALIDADES)),
        "Tipo de Informe": st.column_config.SelectboxColumn(options=opciones(df, "Tipo de Informe", TIPOS)),
        "Estado": st.column_config.SelectboxColumn(options=ESTADOS),
        "Causa de Retraso": st.column_config.SelectboxColumn(options=[""] + opciones(df, "Causa de Retraso", CAUSAS)),
        "N° Revisiones": st.column_config.NumberColumn(min_value=0, step=1),
        **{c: st.column_config.DateColumn(c, format="DD/MM/YYYY") for c in DATE_COLS},
    }
    editado = st.data_editor(df, num_rows="dynamic", use_container_width=True, hide_index=True,
                             column_config=cfg, key="editor_matriz", height=520)
    c1, c2, _ = st.columns([1, 1, 3])
    if c1.button("💾 Guardar cambios", type="primary", key="btn_guardar_editor"):
        guardar(editado, "Matriz actualizada desde el editor")
        st.rerun()
    if c2.button("↩️ Descartar cambios", key="btn_descartar"):
        st.session_state.pop("editor_matriz", None)
        st.rerun()


def puede_editar() -> bool:
    return (not _secret("edicion", "password")) or st.session_state.get("editor_ok", False)


# ════════════════════════════════════════════════════════════════════
# ESTADO INICIAL
# ════════════════════════════════════════════════════════════════════
if "df" not in st.session_state:
    st.session_state.df, st.session_state.origen = cargar_inicial()

df_base: pd.DataFrame = st.session_state.df
HOY = hoy_lima()

# ════════════════════════════════════════════════════════════════════
# BARRA LATERAL: PARÁMETROS, FILTROS Y ARCHIVO
# ════════════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown("### ⚙️ Parámetros del KPI")
    meta = st.slider("Meta de cumplimiento de plazo (OTD)", 50, 100, 90, 1, format="%d%%")
    plazo = st.number_input("Plazo estándar de entrega (días)", 1, 60, 5,
                            help="Se usa cuando el registro no tiene Fecha Compromiso.")
    habiles = st.toggle("Contar en días hábiles (excluye fines de semana y feriados Perú)", value=True)

    datos = calcular(df_base, int(plazo), habiles, HOY)

    st.markdown("### 🔎 Filtros")
    fechas_ok = datos["Fecha Ejecución"].dropna()
    if len(fechas_ok):
        fmin, fmax = fechas_ok.min().date(), max(fechas_ok.max().date(), HOY)
        rango = st.date_input("Rango de Fecha Ejecución", (fmin, fmax), min_value=fmin, max_value=fmax,
                              format="DD/MM/YYYY")
    else:
        rango = ()
    f_esp = st.multiselect("Especialidad", sorted(datos["Especialidad"].unique()))
    f_resp = st.multiselect("Responsable", sorted(datos["Responsable"].unique()))
    f_uni = st.multiselect("Unidad / Área", sorted(datos["Unidad / Área"].unique()))
    f_est = st.multiselect("Estado", [e for e in ESTADOS if e in datos["Estado"].unique()])

    d = datos.copy()
    if isinstance(rango, (tuple, list)) and len(rango) == 2:
        ini, fin = pd.Timestamp(rango[0]), pd.Timestamp(rango[1])
        d = d[d["Fecha Ejecución"].isna() | d["Fecha Ejecución"].between(ini, fin)]
    for col, sel in [("Especialidad", f_esp), ("Responsable", f_resp), ("Unidad / Área", f_uni), ("Estado", f_est)]:
        if sel:
            d = d[d[col].isin(sel)]

    st.markdown("### 📁 Matriz (archivo)")
    if _secret("edicion", "password"):
        if not st.session_state.get("editor_ok"):
            clave = st.text_input("Clave de edición", type="password")
            if clave and clave == _secret("edicion", "password"):
                st.session_state.editor_ok = True
                st.rerun()
            elif clave:
                st.error("Clave incorrecta")
        else:
            st.success("🔓 Edición habilitada")
    ts = st.session_state.get("ultimo_guardado")
    st.caption(f"Registros en matriz: **{len(df_base)}**" + (f" · Último guardado: {ts:%d/%m/%Y %H:%M}" if ts else ""))
    st.download_button("⬇️ Descargar matriz (.xlsx)", a_excel(df_base),
                       file_name=f"MATRIZ_KPI_ENTREGA_INFORMES_{HOY:%Y%m%d}.xlsx", use_container_width=True,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with st.expander("⬆️ Subir matriz actualizada"):
        carga_matriz("sb", compacto=True)

# ════════════════════════════════════════════════════════════════════
# ENCABEZADO
# ════════════════════════════════════════════════════════════════════
st.markdown(
    f"""<div class="hdr"><h1>📊 KPI · Entrega de Informes</h1>
    <p>ADEMINSAC · Control de cumplimiento, tiempos de ciclo y mejora continua · Corte al {HOY:%d/%m/%Y}
    · {'días hábiles' if habiles else 'días calendario'} · plazo estándar {plazo} d · meta OTD {meta}%</p></div>""",
    unsafe_allow_html=True,
)
if "flash" in st.session_state:
    st.success(st.session_state.pop("flash"))
if st.session_state.origen == "demo":
    st.warning("Se muestran **datos de demostración**. Suba su matriz desde la barra lateral o la pestaña "
               "**Gestión de Matriz** para ver sus indicadores reales.", icon="🧪")

ent = d[d["A Tiempo"].notna()]
pend = d[~d["Entregado"]]

tabs = st.tabs(["📊 Resumen Ejecutivo", "🎯 Cumplimiento", "⏱️ Tiempos y Estabilidad",
                "🔧 Mejora Continua", "📋 Backlog y Alertas", "🗂️ Gestión de Matriz"])


def expander_registro(prefix: str) -> None:
    st.divider()
    with st.expander("➕ Añadir registro a la matriz"):
        form_registro(prefix)


# ─────────────────────────── 1. RESUMEN ───────────────────────────
with tabs[0]:
    otd = ent["A Tiempo"].mean() * 100 if len(ent) else np.nan
    lt = d["Lead Time"].mean()
    fpy = d["Sin Revisión"].mean() * 100
    venc = int((pend["Días Vencido"] > 0).sum())
    fmt = lambda v, s="": "—" if pd.isna(v) else f"{v:,.1f}{s}"  # noqa: E731

    c = st.columns(4)
    c[0].metric("Informes registrados", f"{len(d):,}")
    c[1].metric("Entregados", f"{int(d['Entregado'].sum()):,}",
                f"{d['Entregado'].mean() * 100:.0f}% del total" if len(d) else None, delta_color="off")
    c[2].metric("Cumplimiento de plazo (OTD)", fmt(otd, "%"),
                None if pd.isna(otd) else f"{otd - meta:+.1f} pp vs meta")
    c[3].metric("Lead time promedio", fmt(lt, " d"),
                None if pd.isna(lt) else f"{lt - plazo:+.1f} d vs plazo", delta_color="inverse")
    c = st.columns(4)
    c[0].metric("Pendientes", f"{len(pend):,}")
    c[1].metric("Pendientes vencidos", f"{venc:,}",
                f"{venc / len(pend) * 100:.0f}% del backlog" if len(pend) else None, delta_color="off")
    c[2].metric("Aprobados sin revisión (FPY)", fmt(fpy, "%"))
    c[3].metric("Desviación media vs plazo", fmt(ent["Desviación"].mean(), " d"),
                help="Negativo = entregado antes del plazo; positivo = días de atraso.")

    g1, g2 = st.columns([1, 2])
    with g1:
        if pd.notna(otd):
            fig = go.Figure(go.Indicator(
                mode="gauge+number+delta", value=otd,
                number={"suffix": "%", "valueformat": ".1f"},
                delta={"reference": meta, "valueformat": ".1f", "suffix": " pp"},
                gauge={"axis": {"range": [0, 100]}, "bar": {"color": AZUL},
                       "steps": [{"range": [0, max(0, meta - 15)], "color": "#FADBD8"},
                                 {"range": [max(0, meta - 15), meta], "color": "#FDEBD0"},
                                 {"range": [meta, 100], "color": "#D6EAF8"}],
                       "threshold": {"line": {"color": NARANJA, "width": 4}, "value": meta}}))
            mostrar(estilo(fig, 330, "Semáforo de cumplimiento (OTD)"))
        else:
            sin_datos()
    with g2:
        if len(ent):
            m = (ent.groupby("Mes Entrega").agg(Total=("ID", "count"), ATiempo=("A Tiempo", "sum"))
                 .reset_index().sort_values("Mes Entrega"))
            m["Fuera"] = m["Total"] - m["ATiempo"]
            m["OTD"] = m["ATiempo"] / m["Total"] * 100
            fig = make_subplots(specs=[[{"secondary_y": True}]])
            fig.add_bar(x=m["Mes Entrega"], y=m["ATiempo"], name="A tiempo", marker_color=AZUL)
            fig.add_bar(x=m["Mes Entrega"], y=m["Fuera"], name="Fuera de plazo", marker_color=NARANJA)
            fig.add_scatter(x=m["Mes Entrega"], y=m["OTD"], name="% OTD", mode="lines+markers+text",
                            text=m["OTD"].round(0).astype(int).astype(str) + "%", textposition="top center",
                            line=dict(color="#17202A", width=2.5), secondary_y=True)
            fig.add_scatter(x=m["Mes Entrega"], y=[meta] * len(m), name=f"Meta {meta}%", mode="lines",
                            line=dict(color=NARANJA, dash="dash"), secondary_y=True)
            fig.update_layout(barmode="stack")
            fig.update_yaxes(title_text="N° informes", secondary_y=False)
            fig.update_yaxes(title_text="% OTD", range=[0, 110], secondary_y=True, showgrid=False)
            mostrar(estilo(fig, 330, "Entregas mensuales y tendencia del OTD"))
        else:
            sin_datos()

    g3, g4 = st.columns([1, 1])
    with g3:
        s = d["Situación"].value_counts().reindex(list(COLOR_SIT)).dropna().reset_index()
        s.columns = ["Situación", "N"]
        if len(s):
            fig = px.bar(s, x="N", y="Situación", orientation="h", color="Situación",
                         color_discrete_map=COLOR_SIT, text="N")
            fig.update_layout(showlegend=False, yaxis_title=None, xaxis_title="N° informes")
            mostrar(estilo(fig, 320, "Situación actual de la cartera de informes"))
    with g4:
        st.markdown("##### 🧠 Hallazgos automáticos")
        for h in hallazgos(d, meta):
            st.markdown(f"- {h}")
    expander_registro("t1")

# ───────────────────────── 2. CUMPLIMIENTO ─────────────────────────
with tabs[1]:
    if not len(ent):
        sin_datos()
    else:
        g1, g2 = st.columns(2)
        for col, contenedor, titulo in [("Especialidad", g1, "OTD por especialidad"),
                                        ("Responsable", g2, "OTD por responsable")]:
            with contenedor:
                r = (ent.groupby(col)["A Tiempo"].agg(OTD="mean", N="count").reset_index())
                r["OTD"] *= 100
                r = r.sort_values("OTD")
                r["Color"] = np.where(r["OTD"] >= meta, "Cumple meta", "Bajo meta")
                fig = px.bar(r, x="OTD", y=col, orientation="h", color="Color",
                             color_discrete_map={"Cumple meta": AZUL, "Bajo meta": NARANJA},
                             text=r["OTD"].round(0).astype(int).astype(str) + "% (n=" + r["N"].astype(str) + ")",
                             custom_data=["N"])
                fig.add_vline(x=meta, line_dash="dash", line_color="#17202A",
                              annotation_text=f"Meta {meta}%", annotation_position="top")
                fig.update_layout(xaxis_range=[0, 115], xaxis_title="% entregado a tiempo", yaxis_title=None)
                mostrar(estilo(fig, max(320, 42 * len(r) + 80), titulo))

        piv = ent.pivot_table(index="Especialidad", columns="Mes Entrega", values="A Tiempo", aggfunc="mean") * 100
        if piv.shape[1] >= 1:
            fig = px.imshow(piv.sort_index(axis=1), text_auto=".0f", aspect="auto", zmin=0, zmax=100,
                            color_continuous_scale="RdYlBu", labels=dict(color="% OTD", x="Mes", y=""))
            mostrar(estilo(fig, max(300, 50 * len(piv) + 100), "Mapa de calor: OTD por especialidad y mes"))

        u = ent.groupby(["Unidad / Área", "Situación"]).size().reset_index(name="N")
        top = ent["Unidad / Área"].value_counts().head(12).index
        u = u[u["Unidad / Área"].isin(top)]
        u["%"] = u["N"] / u.groupby("Unidad / Área")["N"].transform("sum") * 100
        fig = px.bar(u, x="%", y="Unidad / Área", color="Situación", orientation="h",
                     color_discrete_map=COLOR_SIT, text=u["%"].round(0).astype(int).astype(str) + "%",
                     hover_data=["N"])
        fig.update_layout(barmode="stack", xaxis_title="% de informes entregados", yaxis_title=None)
        mostrar(estilo(fig, max(320, 40 * len(top) + 100), "Composición de cumplimiento por unidad / área (top 12)"))
    expander_registro("t2")

# ────────────────────── 3. TIEMPOS Y ESTABILIDAD ──────────────────────
with tabs[2]:
    lt_df = d[d["Lead Time"].notna()]
    if not len(lt_df):
        sin_datos()
    else:
        g1, g2 = st.columns(2)
        with g1:
            orden = lt_df.groupby("Especialidad")["Lead Time"].median().sort_values().index.tolist()
            fig = px.box(lt_df, x="Especialidad", y="Lead Time", points="outliers",
                         category_orders={"Especialidad": orden}, color_discrete_sequence=[AZUL])
            fig.add_hline(y=plazo, line_dash="dash", line_color=NARANJA, annotation_text=f"Plazo {plazo} d")
            fig.update_layout(yaxis_title="Lead time (días)", xaxis_title=None)
            mostrar(estilo(fig, 360, "Dispersión del lead time por especialidad"))
        with g2:
            fig = px.histogram(lt_df, x="Lead Time", nbins=int(min(40, lt_df["Lead Time"].max() + 1)),
                               color_discrete_sequence=[AZUL_CL])
            fig.add_vline(x=plazo, line_dash="dash", line_color=NARANJA, annotation_text=f"Plazo {plazo} d")
            fig.add_vline(x=lt_df["Lead Time"].mean(), line_color=AZUL,
                          annotation_text=f"Media {lt_df['Lead Time'].mean():.1f} d", annotation_position="top left")
            fig.update_layout(xaxis_title="Lead time (días)", yaxis_title="N° informes", bargap=.05)
            mostrar(estilo(fig, 360, "Distribución del lead time"))

        st.markdown("##### 📉 Gráfico de control I-MR del lead time")
        opts = ["Todas"] + sorted(lt_df["Especialidad"].unique())
        sel = st.selectbox("Especialidad para el gráfico de control", opts, key="cc_esp")
        s = lt_df if sel == "Todas" else lt_df[lt_df["Especialidad"] == sel]
        s = s.sort_values("Fecha Entrega").reset_index(drop=True)
        if len(s) >= 8:
            x = s["Lead Time"].to_numpy()
            media, mrbar = x.mean(), np.abs(np.diff(x)).mean()
            ucl, lcl = media + 2.66 * mrbar, max(0.0, media - 2.66 * mrbar)
            fuera = (x > ucl) | (x < lcl)
            fig = go.Figure()
            fig.add_scatter(x=s.index + 1, y=x, mode="lines+markers", name="Lead time",
                            line=dict(color=AZUL_CL), marker=dict(color=AZUL, size=6),
                            customdata=np.stack([s["ID"], s["Especialidad"], s["Fecha Entrega"].dt.strftime("%d/%m/%Y")], -1),
                            hovertemplate="%{customdata[0]} · %{customdata[1]}<br>Entrega %{customdata[2]}"
                                          "<br>Lead time %{y} d<extra></extra>")
            fig.add_scatter(x=(s.index + 1)[fuera], y=x[fuera], mode="markers", name="Fuera de control",
                            marker=dict(color=ROJO, size=11, symbol="diamond"))
            for y, nom, col, dash in [(media, "Media", "#17202A", "solid"), (ucl, "LSC", NARANJA, "dash"),
                                      (lcl, "LIC", NARANJA, "dash")]:
                fig.add_hline(y=y, line_color=col, line_dash=dash, annotation_text=f"{nom} {y:.1f}",
                              annotation_position="right")
            fig.update_layout(xaxis_title="Secuencia de informes (orden de entrega)", yaxis_title="Días")
            mostrar(estilo(fig, 380))
            st.caption(f"Límites 3σ (I-MR): media {media:.1f} d · LSC {ucl:.1f} d · LIC {lcl:.1f} d · "
                       f"**{int(fuera.sum())}** punto(s) fuera de control por causa especial. "
                       "Si la media está por encima del plazo con el proceso estable, el problema es de "
                       "capacidad del proceso (mejora estructural), no de casos aislados.")
        else:
            sin_datos("Se requieren al menos 8 informes entregados para construir el gráfico de control.")

        t = (lt_df.groupby(["Mes Entrega", "Especialidad"])["Lead Time"].mean().reset_index())
        fig = px.line(t, x="Mes Entrega", y="Lead Time", color="Especialidad", markers=True,
                      color_discrete_sequence=px.colors.qualitative.Safe)
        fig.add_hline(y=plazo, line_dash="dash", line_color=NARANJA, annotation_text=f"Plazo {plazo} d")
        fig.update_layout(yaxis_title="Lead time promedio (días)", xaxis_title=None)
        mostrar(estilo(fig, 360, "Evolución mensual del lead time por especialidad"))
    expander_registro("t3")

# ───────────────────────── 4. MEJORA CONTINUA ─────────────────────────
with tabs[3]:
    g1, g2 = st.columns([3, 2])
    with g1:
        tarde = ent[ent["A Tiempo"] == 0]
        if len(tarde):
            p = (tarde["Causa de Retraso"].replace("", "Sin causa registrada").value_counts()
                 .rename_axis("Causa").reset_index(name="N"))
            p["Acum"] = p["N"].cumsum() / p["N"].sum() * 100
            vital = p["Acum"].shift(fill_value=0) < 80
            fig = make_subplots(specs=[[{"secondary_y": True}]])
            fig.add_bar(x=p["Causa"], y=p["N"], name="Informes fuera de plazo", text=p["N"],
                        marker_color=np.where(vital, NARANJA, GRIS))
            fig.add_scatter(x=p["Causa"], y=p["Acum"], name="% acumulado", mode="lines+markers",
                            line=dict(color=AZUL, width=2.5), secondary_y=True)
            fig.add_scatter(x=p["Causa"], y=[80] * len(p), name="80%", mode="lines",
                            line=dict(color="#17202A", dash="dot"), secondary_y=True)
            fig.update_yaxes(title_text="N° informes", secondary_y=False)
            fig.update_yaxes(title_text="% acumulado", range=[0, 105], secondary_y=True, showgrid=False)
            mostrar(estilo(fig, 400, "Pareto de causas de retraso (80/20)"))
            st.caption("En naranja, las **pocas causas vitales** que explican ~80% de los retrasos: "
                       "son el foco prioritario para el plan de acción (5 Porqués / Ishikawa).")
        else:
            sin_datos("No hay informes fuera de plazo en el periodo filtrado. 👏")
    with g2:
        f = d[d["Sin Revisión"].notna()].groupby("Especialidad")["Sin Revisión"].agg(FPY="mean", N="count").reset_index()
        if len(f):
            f["FPY"] *= 100
            f = f.sort_values("FPY")
            fig = px.bar(f, x="FPY", y="Especialidad", orientation="h", color_discrete_sequence=[AZUL],
                         text=f["FPY"].round(0).astype(int).astype(str) + "%")
            fig.update_layout(xaxis_range=[0, 110], xaxis_title="% aprobado sin revisiones", yaxis_title=None)
            mostrar(estilo(fig, 400, "Calidad a la primera (FPY) por especialidad"))

    g3, g4 = st.columns(2)
    with g3:
        rv = d[d["Entregado"]].copy()
        if len(rv):
            rv["Revisiones"] = pd.cut(rv["N° Revisiones"], [-1, 0, 1, 2, np.inf], labels=["0", "1", "2", "3+"]).astype(str)
            r = rv.groupby(["Especialidad", "Revisiones"]).size().reset_index(name="N")
            fig = px.bar(r, x="Especialidad", y="N", color="Revisiones", text="N",
                         color_discrete_map={"0": AZUL, "1": AZUL_CL, "2": NARANJA, "3+": ROJO},
                         category_orders={"Revisiones": ["0", "1", "2", "3+"]})
            fig.update_layout(barmode="stack", yaxis_title="N° informes", xaxis_title=None)
            mostrar(estilo(fig, 360, "Reprocesos: N° de revisiones por especialidad"))
    with g4:
        if len(ent):
            sc = ent.copy()
            fig = px.scatter(sc, x="N° Revisiones", y="Lead Time", color="Especialidad", opacity=.7,
                             color_discrete_sequence=px.colors.qualitative.Safe,
                             hover_data=["ID", "N° OT", "Responsable"])
            fig.add_hline(y=plazo, line_dash="dash", line_color=NARANJA)
            fig.update_layout(xaxis_title="N° de revisiones", yaxis_title="Lead time (días)")
            mostrar(estilo(fig, 360, "Relación reprocesos vs lead time"))
            if sc["N° Revisiones"].nunique() > 1:
                r = sc[["N° Revisiones", "Lead Time"]].corr().iloc[0, 1]
                st.caption(f"Correlación revisiones–lead time: **r = {r:.2f}**. "
                           + ("Los reprocesos explican parte relevante del atraso: reforzar la revisión "
                              "interna previa a la emisión." if r >= .3 else
                              "Los reprocesos no son el principal impulsor del atraso."))
    expander_registro("t4")

# ─────────────────────── 5. BACKLOG Y ALERTAS ───────────────────────
with tabs[4]:
    if not len(pend):
        st.success("No hay informes pendientes con los filtros actuales. ✅")
    else:
        g1, g2 = st.columns(2)
        with g1:
            b = pend.copy()
            b["Antigüedad (días)"] = pd.cut(b["Antigüedad"], [-np.inf, 3, 7, 15, np.inf],
                                            labels=["0–3", "4–7", "8–15", ">15"]).astype(str)
            a = b.groupby(["Antigüedad (días)", "Especialidad"]).size().reset_index(name="N")
            fig = px.bar(a, x="Antigüedad (días)", y="N", color="Especialidad", text="N",
                         category_orders={"Antigüedad (días)": ["0–3", "4–7", "8–15", ">15"]},
                         color_discrete_sequence=px.colors.qualitative.Safe)
            fig.update_layout(barmode="stack", yaxis_title="N° pendientes")
            mostrar(estilo(fig, 360, "Antigüedad del backlog (aging)"))
        with g2:
            cr = pend.groupby(["Responsable", "Situación"]).size().reset_index(name="N")
            orden = pend["Responsable"].value_counts().sort_values().index.tolist()
            fig = px.bar(cr, x="N", y="Responsable", color="Situación", orientation="h", text="N",
                         color_discrete_map=COLOR_SIT, category_orders={"Responsable": orden[::-1]})
            fig.update_layout(barmode="stack", xaxis_title="N° pendientes", yaxis_title=None)
            mostrar(estilo(fig, max(360, 34 * len(orden) + 100), "Carga pendiente por responsable"))

        st.markdown("##### 🚨 Informes pendientes vencidos")
        v = pend[pend["Días Vencido"] > 0].sort_values("Días Vencido", ascending=False)
        if len(v):
            st.dataframe(
                v[["ID", "N° OT", "Especialidad", "Unidad / Área", "Equipo / TAG", "Responsable",
                   "Fecha Ejecución", "Fecha Límite", "Días Vencido", "Estado"]],
                use_container_width=True, hide_index=True,
                column_config={
                    "Fecha Ejecución": st.column_config.DateColumn(format="DD/MM/YYYY"),
                    "Fecha Límite": st.column_config.DateColumn(format="DD/MM/YYYY"),
                    "Días Vencido": st.column_config.ProgressColumn(
                        format="%d d", min_value=0, max_value=float(v["Días Vencido"].max())),
                })
        else:
            st.success("Todos los pendientes se encuentran dentro del plazo.")
        prox = pend[(pend["Días Vencido"] == 0) & pend["Fecha Límite"].notna()]
        prox = prox[dias_entre(pd.Series(pd.Timestamp(HOY), index=prox.index), prox["Fecha Límite"], habiles) <= 2]
        if len(prox):
            st.markdown("##### ⏳ Por vencer (≤ 2 días)")
            st.dataframe(prox[["ID", "N° OT", "Especialidad", "Responsable", "Fecha Límite", "Estado"]],
                         use_container_width=True, hide_index=True,
                         column_config={"Fecha Límite": st.column_config.DateColumn(format="DD/MM/YYYY")})
    expander_registro("t5")

# ─────────────────────── 6. GESTIÓN DE MATRIZ ───────────────────────
with tabs[5]:
    s1, s2, s3, s4 = st.tabs(["➕ Añadir registro", "✏️ Actualizar matriz", "⬆️ Subir matriz", "⬇️ Descargar"])
    with s1:
        form_registro("gm")
    with s2:
        editor_matriz()
    with s3:
        st.markdown("Suba la matriz actualizada. Las columnas se reconocen automáticamente aunque tengan "
                    "nombres distintos (p. ej. *Fecha Inspección*, *Fecha Límite*, *Disciplina*); puede "
                    "ajustar el mapeo antes de confirmar.")
        carga_matriz("gm")
    with s4:
        kpi_mes = pd.DataFrame()
        if len(ent):
            kpi_mes = (ent.groupby("Mes Entrega").agg(Entregados=("ID", "count"), A_Tiempo=("A Tiempo", "sum"),
                                                      Lead_Time_Prom=("Lead Time", "mean"),
                                                      FPY=("Sin Revisión", "mean")).reset_index())
            kpi_mes["OTD_%"] = (kpi_mes["A_Tiempo"] / kpi_mes["Entregados"] * 100).round(1)
            kpi_mes["FPY"] = (kpi_mes["FPY"] * 100).round(1)
            kpi_mes["Lead_Time_Prom"] = kpi_mes["Lead_Time_Prom"].round(1)
        c1, c2, c3 = st.columns(3)
        c1.download_button("⬇️ Matriz completa (.xlsx)", a_excel(df_base),
                           file_name=f"MATRIZ_KPI_ENTREGA_INFORMES_{HOY:%Y%m%d}.xlsx", use_container_width=True)
        vista = d.drop(columns=["Entregado"]).copy()
        c2.download_button("⬇️ Vista filtrada con cálculos + KPI mensual", a_excel(vista, {"KPI_MENSUAL": kpi_mes}),
                           file_name=f"REPORTE_KPI_INFORMES_{HOY:%Y%m%d}.xlsx", use_container_width=True)
        c3.download_button("⬇️ Plantilla vacía", a_excel(pd.DataFrame(columns=COLUMNS)),
                           file_name="PLANTILLA_MATRIZ_KPI_INFORMES.xlsx", use_container_width=True)
        if len(kpi_mes):
            st.dataframe(kpi_mes, use_container_width=True, hide_index=True)

st.caption("OTD: entregado ≤ fecha límite (Fecha Compromiso o, en su defecto, Fecha Ejecución + plazo estándar). "
           "Lead time: Fecha Ejecución → Fecha Entrega. FPY: informes entregados con 0 revisiones. "
           "Los registros «Anulado» se excluyen de los indicadores.")
